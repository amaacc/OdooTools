import os
import sys
import re
import math
import json
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from typing import Dict, List, Optional, Tuple, Any, Callable

from odoo_verifier import OdooVerifier

DEFAULT_ODOO_CATEGORIES = [
    "All",
    "All / Expenses",
    "GCP-General Cosmetics Products",
    "GCS-General Cosmetics Samples",
    "GCT-General Cosmetics Testers",
    "GIP-General Imported Products",
    "GP-GENERAL PRODUCT",
    "GPS-GENERAL PRODUCT SAMPLE",
    "GPT-GENERAL PRODUCT TESTER",
    "HCP-Hair Care Products",
    "HCS-Hair Care Samples",
    "HCT-Hair Care Testers",
    "HHP-Houshold Products",
    "HHS-Houshold Samples",
    "HHT-Houshold Testers",
    "HQP-High Quality Perfumes",
    "HQS-High Quality Samples",
    "HQT-High Quality Testers",
    "MAP-Mass Quality Perfumes",
    "MAS-Mass Quality Samples",
    "MAT-Mass Quality Testers",
    "MDP-Medium Quality Perfumes",
    "MDS-Medium Quality Samples",
    "MDT-Medium Quality Testers",
    "PSP-Personal Care Products",
    "PSS-Personal Care Samples",
    "PST-Personal Care Testers",
    "SCP-Skin Care Products",
    "SCS-Skin Care",
    "SCS-Skin Care Samples",
    "SCT-Skin Care Testers",
    "SET-GIFT COLLECT",
    "STF-Display Stand"
]

DEFAULT_TAG_SUFFIX_RULES = {
    "SG0040": "8",
    "SG0010": "4",
    "SG0060": "4",
    "SG0020": "3",
    "SG0030": "3"
}

DEFAULT_CATEGORY_PREFIX_RULES = {
    "All": "4606-",
    "All / Expenses": "4606-",
    "GCP-General Cosmetics Products": "4303-",
    "GCS-General Cosmetics Samples": "4303-",
    "GCT-General Cosmetics Testers": "4303-",
    "GIP-General Imported Products": "4505-",
    "GP-GENERAL PRODUCT": "4606-",
    "GPS-GENERAL PRODUCT SAMPLE": "4606-",
    "GPT-GENERAL PRODUCT TESTER": "4606-",
    "HCP-Hair Care Products": "4505-",
    "HCS-Hair Care Samples": "4505-",
    "HCT-Hair Care Testers": "4505-",
    "HHP-Houshold Products": "4404-",
    "HHS-Houshold Samples": "4404-",
    "HHT-Houshold Testers": "4404-",
    "HQP-High Quality Perfumes": "4101-",
    "HQS-High Quality Samples": "4101-",
    "HQT-High Quality Testers": "4101-",
    "MAP-Mass Quality Perfumes": "4202-",
    "MAS-Mass Quality Samples": "4202-",
    "MAT-Mass Quality Testers": "4202-",
    "MDP-Medium Quality Perfumes": "4202-",
    "MDS-Medium Quality Samples": "4202-",
    "MDT-Medium Quality Testers": "4202-",
    "PSP-Personal Care Products": "4505-",
    "PSS-Personal Care Samples": "4505-",
    "PST-Personal Care Testers": "4505-",
    "SCP-Skin Care Products": "4505-",
    "SCS-Skin Care": "4505-",
    "SCS-Skin Care Samples": "4505-",
    "SCT-Skin Care Testers": "4505-",
    "SET-GIFT COLLECT": "4101-",
    "STF-Display Stand": "4606-"
}

KEYWORD_PREFIX_MAP = {
    "HQP": "4101-", "HQS": "4101-", "HQT": "4101-", "GIFT": "4101-", "HIGH QUALITY": "4101-",
    "MAP": "4202-", "MAS": "4202-", "MAT": "4202-", "MASS": "4202-",
    "MDP": "4202-", "MDS": "4202-", "MDT": "4202-", "MEDIUM": "4202-",
    "GCP": "4303-", "GCS": "4303-", "GCT": "4303-", "COSMETIC": "4303-", "تجميل": "4303-",
    "HHP": "4404-", "HHS": "4404-", "HHT": "4404-", "HOUSEHOLD": "4404-", "منزلي": "4404-",
    "HCP": "4505-", "HCS": "4505-", "HCT": "4505-", "HAIR": "4505-", "شعر": "4505-",
    "SCP": "4505-", "SCS": "4505-", "SCT": "4505-", "SKIN": "4505-", "بشرة": "4505-",
    "PSP": "4505-", "PSS": "4505-", "PST": "4505-", "PERSONAL": "4505-", "عناية": "4505-",
    "GIP": "4505-", "IMPORTED": "4505-", "مستورد": "4505-",
    "GP": "4606-", "GPS": "4606-", "GPT": "4606-", "STF": "4606-", "EXPENSES": "4606-", "ALL": "4606-"
}

DEFAULT_TAX_RULES = {
    "taxable_keywords": [
        "15", "yes", "نعم", "خاضع", "خاضعة", "vat", "tax", "taxable",
        "شامل", "مضافة", "مفروض", "مفروضة", "standard", "ضريبة", "قيمة مضافة"
    ],
    "exempt_keywords": [
        "معف", "غير خاضع", "غير خاضعة", "بدون", "exempt", "zero", "0%", "لا", "no"
    ]
}


def parse_tax_status(val: Any, custom_rules: Optional[Dict[str, Any]] = None) -> Tuple[bool, bool]:
    """
    Parses tax value from supplier file into (is_taxable, is_valid).
    Supports:
    - Boolean: True/False
    - Numbers: 0 -> (False, True), 0.15/15 -> (True, True)
    - Compound phrases: "Yes - 15% VAT", "15% VAT", "VAT 15%", "0.15", "15%"
    - Keywords: "نعم", "خاضع", "شامل", "معفى", "غير خاضع", "لا", "exempt", "zero"
    - Empty / None: Defaults to (True, True) [Saudi standard retail VAT]
    """
    if val is None:
        return (True, True)
    if isinstance(val, bool):
        return (val, True)
    if isinstance(val, (int, float)):
        if abs(float(val)) < 1e-6:
            return (False, True)
        elif float(val) > 0:
            return (True, True)
        return (False, False)

    s = str(val).strip().lower()
    if not s:
        return (True, True)

    s_clean = re.sub(r'\s+', ' ', s)

    # 1. Check numeric tokens (e.g. 15, 0.15, 0, 0.0)
    numbers = re.findall(r'(\d+(?:\.\d+)?)', s_clean)
    has_15 = False
    has_zero = False
    for n in numbers:
        try:
            fn = float(n)
            if abs(fn - 15.0) < 1e-4 or abs(fn - 0.15) < 1e-4:
                has_15 = True
            elif abs(fn - 0.0) < 1e-4:
                has_zero = True
        except ValueError:
            pass

    taxable_kw = list(DEFAULT_TAX_RULES["taxable_keywords"])
    exempt_kw = list(DEFAULT_TAX_RULES["exempt_keywords"])

    if custom_rules:
        if "taxable_keywords" in custom_rules and isinstance(custom_rules["taxable_keywords"], list):
            taxable_kw = [str(k).strip().lower() for k in custom_rules["taxable_keywords"] if str(k).strip()]
        if "exempt_keywords" in custom_rules and isinstance(custom_rules["exempt_keywords"], list):
            exempt_kw = [str(k).strip().lower() for k in custom_rules["exempt_keywords"] if str(k).strip()]

    if has_15:
        return (True, True)

    if has_zero and not any(kw in s_clean for kw in ['yes', 'نعم', 'خاضع', 'vat', 'tax', '15']):
        return (False, True)

    # Check exempt keywords with boundary check for short words
    for ekw in exempt_kw:
        if len(ekw) <= 2:
            if re.search(r'(?:^|[^\w\u0600-\u06FF])' + re.escape(ekw) + r'(?:[^\w\u0600-\u06FF]|$)', s_clean):
                return (False, True)
        else:
            if ekw in s_clean:
                return (False, True)

    # Check taxable keywords with boundary check for short words
    for tkw in taxable_kw:
        if len(tkw) <= 2:
            if re.search(r'(?:^|[^\w\u0600-\u06FF])' + re.escape(tkw) + r'(?:[^\w\u0600-\u06FF]|$)', s_clean):
                return (True, True)
        else:
            if tkw in s_clean:
                return (True, True)

    return (True, False)




def validate_barcode_checksum(barcode: str) -> bool:
    """
    Validates EAN-13, EAN-8, or UPC-A barcode checksum using GS1 modulo-10 algorithm.
    """
    clean = str(barcode).strip()
    digits = [int(c) for c in clean if c.isdigit()]
    if len(digits) not in (8, 12, 13, 14):
        return False

    check_digit = digits[-1]
    payload = digits[:-1]

    # Weigh digits from right to left: alternating 3 and 1
    total = sum(d * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(payload)))
    calc_check = (10 - (total % 10)) % 10
    return calc_check == check_digit


def calculate_sales_price(cost: float, margin: float = 0.20, is_taxable: bool = True, tax_rate: float = 0.15) -> float:
    """
    Calculates retail sales price:
    - Real profit margin: cost / (1 - margin)
    - Adds VAT if taxable: price * (1 + tax_rate)
    - Rounds up to nearest quarter (0.25)
    - Halala Rule: If result is a whole number (e.g. 100.00), subtracts 1 halala (99.99).
    - If cost <= 0, returns 0.0 (no negative reduction).
    """
    if not cost or cost <= 0:
        return 0.0

    margin = max(0.0, min(0.95, margin))
    price_excl_vat = float(cost) / (1.0 - margin)
    raw_price = price_excl_vat * (1.0 + tax_rate if is_taxable else 1.0)

    val = round(raw_price, 4)
    quarter = math.ceil(val * 4.0) / 4.0

    # Whole number check (subtract 1 halala)
    if abs(quarter - round(quarter)) < 1e-6:
        return round(quarter - 0.01, 2)
    return round(quarter, 2)


COLUMN_SYNONYMS = {
    'ea_barcode': [
        'ea barcode', 'eabarcode', 'ea_barcode', 'باركود الحبة', 'barcode', 'bar code',
        'باركود', 'الباركود', 'item barcode', 'upc', 'ean', 'unit barcode', 'single barcode'
    ],
    'cs_barcode': [
        'cs barcode', 'csbarcode', 'cs_barcode', 'carton barcode', 'باركود الكرتون',
        'case barcode', 'باركود الصندوق', 'outer barcode'
    ],
    'ea_cost': [
        'ea cost (excl vat)', 'ea cost', 'cost (excl vat)', 'تكلفة الحبه قبل الضريبه',
        'تكلفة الحبة', 'تكلفة الشراء', 'cost price', 'purchase price', 'unit cost',
        'سعر التكلفة', 'التكلفة', 'تكلفة', 'cost'
    ],
    'cs_cost': [
        'cs cost (excl vat)', 'cs cost', 'تكلفة الكرتون', 'carton cost', 'case cost'
    ],
    'cs_size': [
        'cs size', 'carton size', 'كميه الكرتون', 'حجم الكرتون', 'عبوة الكرتون',
        'packing', 'pack size', 'سعة الكرتون', 'units per case'
    ],
    'name_en': [
        'productnameen', 'product name en', 'description en', 'english description',
        'item name', 'إسم الصنف انجليزي', 'اسم الصنف انجليزي', 'الوصف بالانجليزي',
        'اسم الصنف بالانجليزي', 'product name', 'item description', 'description'
    ],
    'name_ar': [
        'productnamear', 'product name ar', 'description ar', 'arabic description',
        'اسم الصنف عربي', 'إسم الصنف عربي', 'الوصف بالعربي', 'اسم الصنف بالعربي',
        'اسم المنتج', 'اسم الصنف'
    ],
    'brand': [
        'brand', 'brand name', 'ماركة', 'الماركة', 'براند', 'إسم البراند', 'اسم البراند'
    ],
    'category': [
        'category', 'cat', 'الكاتجري', 'التصنيف', 'الفئة', 'قسم', 'فئة الصنف', 'product category'
    ],
    'sub_category': [
        'sub category', 'sub-category', 'subcategory', 'الصب كاتجري', 'التصنيف الفرعي',
        'الفئة الفرعية', 'sub category raw'
    ],
    'tax': [
        'tax', 'vat', 'ضريبة', 'الضريبة', 'هل المنتج عليه ضريبه او لا', 'taxes', 'tax rate',
        'ضريبة القيمة المضافة', 'قيمة مضافة', 'خاضع للضريبة', 'نسبة الضريبة', 'vat rate', 'vat %'
    ],
    'origin': [
        'country of origin', 'origin', 'بلد التصنيع', 'بلد المنشأ', 'المنشأ', 'صناعة', 'made in'
    ],
    'picture': [
        'picture', 'image', 'photo', 'الصورة', 'رابط الصورة', 'صورة الصنف'
    ],
    'shelf_life': [
        'shelf life (days)', 'shelf life', 'كم سنه صلاحية المنتج', 'فترة الصلاحية',
        'الصلاحية', 'مدة الصلاحية'
    ],
}


def clean_header(val: Any) -> str:
    """Normalizes header strings for flexible matching."""
    if val is None:
        return ''
    s = str(val).strip().lower()
    s = re.sub(r'[\r\n\t]+', ' ', s)
    s = re.sub(r'[\(\)\[\]\{\}\:\-_/\\*]+', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s


def detect_header_row_and_mapping(ws) -> Tuple[int, Dict[str, int], Dict[int, str]]:
    """
    Scans the first 15 rows of the worksheet to dynamically discover:
    1. Header row index (1-indexed)
    2. Column mapping: field_name -> col_idx (0-indexed)
    3. Raw headers detected: col_idx -> original cell value
    """
    best_row = 3
    max_matches = 0
    best_mapping: Dict[str, int] = {}
    best_raw_headers: Dict[int, str] = {}

    for r in range(1, min(ws.max_row + 1, 16)):
        row_cells = ws[r]
        cleaned_headers = [clean_header(cell.value) for cell in row_cells]
        raw_headers = {idx: str(cell.value or '').strip() for idx, cell in enumerate(row_cells)}

        mapping: Dict[str, int] = {}
        assigned_cols = set()

        # Pass 1: Exact matches
        for field, syns in COLUMN_SYNONYMS.items():
            if field in mapping:
                continue
            for syn in syns:
                clean_syn = clean_header(syn)
                for col_idx, h in enumerate(cleaned_headers):
                    if col_idx in assigned_cols or not h:
                        continue
                    if h == clean_syn:
                        mapping[field] = col_idx
                        assigned_cols.add(col_idx)
                        break
                if field in mapping:
                    break

        # Pass 2: Substring matches for remaining fields
        for field, syns in COLUMN_SYNONYMS.items():
            if field in mapping:
                continue
            for syn in syns:
                clean_syn = clean_header(syn)
                for col_idx, h in enumerate(cleaned_headers):
                    if col_idx in assigned_cols or not h:
                        continue
                    if clean_syn in h:
                        mapping[field] = col_idx
                        assigned_cols.add(col_idx)
                        break
                if field in mapping:
                    break

        matches_count = len(mapping)
        has_essential = ('ea_barcode' in mapping) or ('ea_cost' in mapping) or ('name_en' in mapping)
        if has_essential and matches_count >= max_matches and matches_count >= 2:
            max_matches = matches_count
            best_row = r
            best_mapping = mapping
            best_raw_headers = raw_headers

    return best_row, best_mapping, best_raw_headers


def safe_float(val: Any) -> float:
    """
    Safely converts a value to float.
    Handles currency symbols, commas, spaces, percentages, formulas.
    Never throws an exception. Returns 0.0 on failure.
    """
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s:
        return 0.0
    s = re.sub(r'[\$,\sSARsarر\.س]+', '', s)
    if s.endswith('%'):
        s = s[:-1]
    try:
        return float(s)
    except Exception:
        m = re.search(r'[-+]?\d*\.?\d+', s)
        if m:
            try:
                return float(m.group(0))
            except Exception:
                pass
        return 0.0


def extract_vendor_hint(file_path: str, ws: Any = None) -> Dict[str, str]:
    """
    Extracts potential vendor code and name from:
    1. Top rows of Excel (if 'Supplier Name' exists in row 1 or 2)
    2. File name pattern (e.g. '905-315 AlArkan.xlsx' -> '905-315')
    """
    code = ""
    name = ""

    # 1. From Worksheet Row 1 or 2
    if ws is not None:
        try:
            for r in (1, 2):
                if ws.max_row >= r:
                    vals = [cell.value for cell in ws[r]]
                    if len(vals) >= 2 and vals[1]:
                        val1 = str(vals[0] or '').strip().lower()
                        val2 = str(vals[1] or '').strip()
                        if 'supp' in val1 or 'مورد' in val1 or re.match(r'^\d{2,4}[-\/]\d{2,4}$', val2):
                            code = val2
                            if len(vals) >= 3 and vals[2]:
                                name = str(vals[2]).strip()
                            break
        except Exception:
            pass

    # 2. From File Name if not found in sheet
    if not code and file_path:
        base = os.path.splitext(os.path.basename(file_path))[0]
        # Match pattern like 905-315 or 10045
        m = re.search(r'(?:^|[_\s\-])([A-Za-z]?\d{2,4}[-\/]\d{2,4}|\d{3,6})(?:$|[_\s\-])', base)
        if m:
            code = m.group(1).strip()
            remaining = base.replace(m.group(0), ' ').replace('_', ' ').replace('-', ' ')
            rem_words = [w for w in remaining.split() if w.lower() not in ('suppitems', 'items', 'products', 'new', 'item', 'supp')]
            if rem_words:
                name = " ".join(rem_words)
        else:
            tokens = [t for t in re.split(r'[_\-\s]+', base) if t.lower() not in ('suppitems', 'items', 'products', 'new', 'item', 'supp', 'test')]
            if tokens:
                code = tokens[0]
                if len(tokens) > 1:
                    name = " ".join(tokens[1:])

    return {'code': code, 'name': name}


class ImportEngine:

    """
    Main processing engine that orchestrates reading suppItems.xlsx,
    verifying with Odoo, calculating prices, formatting Odooimport.xlsx,
    generating comparison sheets, and updating AllNewItems.xlsx.
    """

    def __init__(self, config_path: str = None):
        if not config_path:
            config_path = self._resolve_config_path()
        self.config_path = config_path
        self.config = self.load_config()
        self.verifier = OdooVerifier()

    @staticmethod
    def _is_valid_importer_config_file(path: str) -> bool:
        if not path or not os.path.exists(path):
            return False
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, dict):
                    return "category_prefix_rules" in data or "internal_reference_ranges" in data or "default_tag" in data.get("settings", {})
        except Exception:
            return False
        return False

    def _resolve_config_path(self) -> str:
        candidates = []
        if getattr(sys, 'frozen', False):
            exe_dir = os.path.dirname(sys.executable)
            candidates.append(os.path.join(exe_dir, "importer_config.json"))
            candidates.append(os.path.join(exe_dir, "config.json"))
            meipass = getattr(sys, '_MEIPASS', None)
            if meipass:
                candidates.append(os.path.join(meipass, "importer_config.json"))
                candidates.append(os.path.join(meipass, "config.json"))
        else:
            base_dir = os.path.dirname(os.path.abspath(__file__))
            candidates.append(os.path.join(base_dir, "importer_config.json"))
            candidates.append(os.path.join(base_dir, "config.json"))

        for c in candidates:
            if self._is_valid_importer_config_file(c):
                return c

        target_dir = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
        return os.path.join(target_dir, "importer_config.json")

    def load_config(self) -> Dict[str, Any]:
        cfg = {
            "settings": {
                "default_profit_margin": 0.20,
                "default_tax_rate": 0.15,
                "default_tag": "SG0040-OTHER PRODUCTS",
                "default_category": "SCP-Skin Care Products",
                "uom_po_id": "PC",
                "uom_so_id": "PC",
                "customer_tax_id": 1,
                "vendor_tax_id": 3,
                "customer_tax_id_zero": 45,
                "vendor_tax_id_zero": 9,
                "type": "Storable Product",
                "invoice_policy": "Delivered quantities",
                "available_in_pos": True,
                "sequence": 1
            },
            "tag_suffix_rules": dict(DEFAULT_TAG_SUFFIX_RULES),
            "category_prefix_rules": dict(DEFAULT_CATEGORY_PREFIX_RULES),
            "tax_rules": dict(DEFAULT_TAX_RULES),
            "internal_reference_ranges": [],
            "category_aliases": {}
        }
        if self.config_path and os.path.exists(self.config_path):
            try:
                with open(self.config_path, "r", encoding="utf-8") as f:
                    loaded = json.load(f)
                    if isinstance(loaded, dict):
                        if "settings" in loaded and isinstance(loaded["settings"], dict):
                            cfg["settings"].update(loaded["settings"])
                        if "tag_suffix_rules" in loaded and isinstance(loaded["tag_suffix_rules"], dict):
                            cfg["tag_suffix_rules"].update(loaded["tag_suffix_rules"])
                        if "category_prefix_rules" in loaded and isinstance(loaded["category_prefix_rules"], dict):
                            cfg["category_prefix_rules"].update(loaded["category_prefix_rules"])
                        if "tax_rules" in loaded and isinstance(loaded["tax_rules"], dict):
                            cfg["tax_rules"] = loaded["tax_rules"]
                        if "internal_reference_ranges" in loaded and isinstance(loaded["internal_reference_ranges"], list) and loaded["internal_reference_ranges"]:
                            cfg["internal_reference_ranges"] = loaded["internal_reference_ranges"]
                        if "category_aliases" in loaded and isinstance(loaded["category_aliases"], dict):
                            cfg["category_aliases"].update(loaded["category_aliases"])
                        return cfg
            except Exception as e:
                print(f"[ImportEngine] Error reading config: {e}")
        return cfg

    def save_config(self):
        try:
            with open(self.config_path, "w", encoding="utf-8") as f:
                json.dump(self.config, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"[ImportEngine] Error saving config: {e}")

    def get_odoo_category_names(self) -> List[str]:
        names = []
        if self.verifier and self.verifier.uid:
            try:
                names = self.verifier.get_category_names_list()
            except Exception:
                pass
        if not names:
            cfg_cats = list(self.config.get("category_prefix_rules", {}).keys())
            if cfg_cats:
                names = sorted(cfg_cats)
            else:
                names = list(DEFAULT_ODOO_CATEGORIES)
        return names

    def find_internal_ref_prefix(self, category_str: str, tag_str: str) -> str:
        """
        Determines the internal reference range prefix based on Category & Tag.
        Category determines the base prefix (e.g. '4505-'), Tag determines suffix digit ('4' or '8').
        Example: SCP + SG0040 -> '4505-8', GCP + SG0020 -> '4303-3'.
        """
        cat_upper = (category_str or "").strip().upper()
        tag_upper = (tag_str or "").strip().upper()

        # 1. Determine Suffix Digit from Tag
        tag_suffix_rules = dict(DEFAULT_TAG_SUFFIX_RULES)
        tag_suffix_rules.update(self.config.get("tag_suffix_rules", {}))

        suffix_digit = ""
        for t_code, digit in tag_suffix_rules.items():
            if t_code.upper() in tag_upper:
                suffix_digit = str(digit).strip()
                break

        if not suffix_digit:
            if "40" in tag_upper:
                suffix_digit = "8"
            elif "20" in tag_upper or "30" in tag_upper:
                suffix_digit = "3"
            elif "10" in tag_upper or "60" in tag_upper:
                suffix_digit = "4"
            else:
                suffix_digit = "8"

        # 2. Determine Base Prefix from Category Rules
        cat_rules = dict(DEFAULT_CATEGORY_PREFIX_RULES)
        cat_rules.update(self.config.get("category_prefix_rules", {}))

        base_prefix = ""

        # Exact match
        for c_name, pre in cat_rules.items():
            if c_name.upper() == cat_upper:
                base_prefix = pre.strip()
                break

        # Substring / Code Match (e.g. HCP or HQP or GCP)
        if not base_prefix:
            m = re.match(r'^([A-Z0-9]+)', cat_upper)
            code_prefix = m.group(1) if m else ""
            if code_prefix:
                for c_name, pre in cat_rules.items():
                    if code_prefix in c_name.upper() or c_name.upper().startswith(code_prefix):
                        base_prefix = pre.strip()
                        break
                if not base_prefix and code_prefix in KEYWORD_PREFIX_MAP:
                    base_prefix = KEYWORD_PREFIX_MAP[code_prefix]

        # Keyword match
        if not base_prefix:
            for kw, pre in KEYWORD_PREFIX_MAP.items():
                if kw in cat_upper:
                    base_prefix = pre
                    break

        if not base_prefix:
            for c_name, pre in cat_rules.items():
                if c_name.upper() in cat_upper or cat_upper in c_name.upper():
                    base_prefix = pre.strip()
                    break

        # 3. Check legacy internal_reference_ranges if any explicit override exists
        ranges = self.config.get("internal_reference_ranges", [])
        for r in ranges:
            cc = r.get("category_code", "").upper()
            tc = r.get("tag_code", "").upper()
            if cc and tc and tc != "ALL" and cc in cat_upper and tc in tag_upper:
                return r.get("prefix", f"{base_prefix or '4505-'}{suffix_digit}")

        if not base_prefix:
            base_prefix = "4505-"

        if not base_prefix.endswith("-"):
            base_prefix += "-"

        return f"{base_prefix}{suffix_digit}"

    def resolve_category(self, raw_category: str, odoo_cats: Dict[str, Any], default_fallback: str = "") -> Tuple[str, bool]:
        """
        Matches raw category with aliases and Odoo categories.
        Returns: (final_category_name, is_valid)
        """
        clean = (raw_category or "").strip()
        if not clean:
            if default_fallback and default_fallback.strip():
                clean = default_fallback.strip()
            else:
                return ("", False)

        aliases = self.config.get("category_aliases", {})
        clean_lower = clean.lower()

        # Check aliases dictionary
        target_name = aliases.get(clean_lower, clean)

        # Check in Odoo categories
        if target_name.lower() in odoo_cats:
            return (odoo_cats[target_name.lower()]['name'], True)

        # Check by prefix code (e.g. SCP)
        code_match = re.match(r'^([A-Z0-9]+)', target_name)
        if code_match and code_match.group(1).lower() in odoo_cats:
            return (odoo_cats[code_match.group(1).lower()]['name'], True)

        # If not found in Odoo, return target_name but flagged invalid
        return (target_name, False)

    def process_file(
        self,
        supp_items_path: str,
        output_import_path: str,
        margin_override: Optional[float] = None,
        vendor_override: Optional[Dict[str, Any]] = None,
        tag_override: Optional[str] = None,
        category_override: Optional[str] = None,
        progress_cb: Optional[Callable[[float, str], None]] = None
    ) -> Dict[str, Any]:
        """
        Reads suppItems.xlsx, validates against rules & Odoo,
        creates Odooimport.xlsx with import sheet and duplicate comparison sheet.
        """
        def notify(pct, msg):
            if progress_cb:
                progress_cb(pct, msg)

        notify(5, "جاري فتح ملف المورد وقراءة البيانات...")

        wb_supp = openpyxl.load_workbook(supp_items_path, data_only=True)
        ws_supp = wb_supp.active

        # 1. Resolve Vendor / Partner (from override, filename, or sheet headers)
        notify(8, "جاري تحديد المورد وفحص الوسوم مع أودو 17...")
        if vendor_override and vendor_override.get('tag_name'):
            vendor_info = vendor_override
        else:
            hint = extract_vendor_hint(supp_items_path, ws_supp)
            query = hint.get('code') or hint.get('name') or ""
            if query:
                vendor_info = self.verifier.resolve_vendor(query)
            else:
                vendor_info = {'success': False}
            if not vendor_info.get('success'):
                s_no = hint.get('code') or "SUPP"
                s_name = hint.get('name') or "Supplier"
                vendor_info = {
                    'success': False,
                    'source': 'hint',
                    'tag_name': f"{s_no} {s_name}".strip(),
                    'supplier_no': s_no,
                    'supplier_name': s_name,
                    'is_tag_in_odoo': False
                }

        supplier_no = vendor_info.get('supplier_no', 'SUPP')
        supplier_name = vendor_info.get('supplier_name', 'Supplier')
        supplier_tag = vendor_info.get('tag_name', f"{supplier_no} {supplier_name}".strip())
        partner_id = vendor_info.get('partner_id')

        # 2. Dynamic Header Row & Column Mapping
        notify(15, "جاري الكشف التلقائي عن ترويسة الأعمدة ومطابقتها...")
        header_row_idx, col_map, raw_headers = detect_header_row_and_mapping(ws_supp)

        if 'ea_barcode' not in col_map and 'name_en' not in col_map:
            detected_cols_str = ", ".join([f"'{v}'" for v in list(raw_headers.values())[:8] if v])
            raise ValueError(
                f"تعذر العثور على عمود الباركود أو اسم الصنف في ملف المورد تلقائياً.\n"
                f"السطر المكتشف: {header_row_idx}\n"
                f"عناوين الأعمدة بالملف: {detected_cols_str}"
            )

        data_rows = []
        barcodes_list = []
        for r in range(header_row_idx + 1, ws_supp.max_row + 1):
            row_cells = ws_supp[r]
            row_vals = [cell.value for cell in row_cells]
            if not any(v is not None for v in row_vals):
                continue

            def get_val(f_name: str, default: Any = ""):
                idx = col_map.get(f_name)
                if idx is not None and idx < len(row_vals):
                    v = row_vals[idx]
                    return v if v is not None else default
                return default

            brand = str(get_val('brand', '')).strip()
            name_en = str(get_val('name_en', '')).strip()
            name_ar = str(get_val('name_ar', '')).strip()
            ea_barcode = str(get_val('ea_barcode', '')).strip()
            cs_barcode = str(get_val('cs_barcode', '')).strip()
            cs_size = get_val('cs_size', None)
            cs_cost = safe_float(get_val('cs_cost', 0.0))
            ea_cost = safe_float(get_val('ea_cost', 0.0))
            origin = str(get_val('origin', '')).strip()
            picture = str(get_val('picture', '')).strip()
            shelf_life = str(get_val('shelf_life', '')).strip()
            tax_raw = get_val('tax', None)
            category_raw = str(get_val('category', '')).strip()
            sub_category_raw = str(get_val('sub_category', '')).strip()

            if not name_en and not ea_barcode:
                continue

            item = {
                'row_idx': r,
                'brand': brand,
                'name_en': name_en,
                'name_ar': name_ar,
                'ea_barcode': ea_barcode,
                'cs_barcode': cs_barcode,
                'cs_size': cs_size,
                'cs_cost': cs_cost,
                'ea_cost': ea_cost,
                'origin': origin,
                'picture': picture,
                'shelf_life': shelf_life,
                'tax_raw': tax_raw,
                'category_raw': category_raw,
                'sub_category_raw': sub_category_raw
            }
            data_rows.append(item)
            if ea_barcode:
                barcodes_list.append(ea_barcode)

        notify(25, f"تم قراءة {len(data_rows)} صنف (سطر الترويسة: {header_row_idx}). جاري فحص البيانات مع أودو 17...")


        # 3. Connect to Odoo and Fetch Metadata
        has_odoo = self.verifier.authenticate()
        existing_products: Dict[str, Dict[str, Any]] = {}
        odoo_cats: Dict[str, Any] = {}
        odoo_tags: Dict[str, Any] = {}

        if has_odoo:
            notify(35, "الاتصال بنجاح مع أودو. جاري فحص الباركودات والتصنيفات والوسوم...")
            odoo_cats = self.verifier.get_all_categories()
            odoo_tags = self.verifier.get_all_tags()
            existing_products = self.verifier.check_existing_products_by_barcodes(barcodes_list)
        else:
            err_msg = self.verifier.last_error or "فشل الاتصال"
            notify(35, f"⚠️ تعذر الاتصال بسيرفر أودو ({err_msg}) - سيتم استخدام القواعد المحلية والفحص الرياضي للباركود.")

        # Settings
        settings = self.config.get("settings", {})
        margin = margin_override if margin_override is not None else settings.get("default_profit_margin", 0.20)
        default_tag = tag_override if tag_override else settings.get("default_tag", "SG0040-OTHER PRODUCTS")
        default_cat = category_override if category_override else settings.get("default_category", "SCP-Skin Care Products")
        uom_po_id = settings.get("uom_po_id", "PC")
        uom_so_id = settings.get("uom_so_id", "PC")
        customer_tax_id = settings.get("customer_tax_id", 1)
        vendor_tax_id = settings.get("vendor_tax_id", 3)
        customer_tax_id_zero = settings.get("customer_tax_id_zero", 45)
        vendor_tax_id_zero = settings.get("vendor_tax_id_zero", 9)
        prod_type = settings.get("type", "Storable Product")
        invoice_policy = settings.get("invoice_policy", "Delivered quantities")
        available_in_pos = settings.get("available_in_pos", True)
        fixed_sequence = settings.get("sequence", 1)

        # Sequence Generators per Prefix
        seq_counters: Dict[str, int] = {}

        def get_next_seq_code(prefix: str) -> str:
            if prefix not in seq_counters:
                if has_odoo:
                    base_seq = self.verifier.get_highest_internal_reference_sequence(prefix)
                else:
                    base_seq = 0
                seq_counters[prefix] = base_seq

            seq_counters[prefix] += 1
            seq = seq_counters[prefix]
            return f"{prefix}{seq:05d}"

        notify(50, "جاري معالجة الأصناف وحساب الأسعار وتوليد الأكواد...")

        # Process Rows
        processed_rows = []
        duplicate_rows = []
        missing_new_tags = {}  # tag_name_upper -> {'name': tag_name_upper, 'count': int, 'sample_barcodes': list}
        total_errors = 0

        for item in data_rows:
            row_errors = []
            row_warnings = []

            # A. Barcode Checksum & Duplicate
            bc = item['ea_barcode']
            is_bc_valid_chk = validate_barcode_checksum(bc) if bc else False
            is_bc_duplicate = bc in existing_products if bc else False

            if not bc:
                row_errors.append("الباركود فارغ")
            elif not is_bc_valid_chk:
                row_errors.append("الباركود غير صالح رياضياً (خطأ في Checksum)")

            if is_bc_duplicate:
                row_warnings.append(f"الباركود مكرر في أودو ({existing_products[bc].get('default_code') or 'مسجل مسبقاً'})")

            # B. Category Resolution
            raw_cat = (item['category_raw'] or "").strip()
            if not raw_cat and default_cat:
                raw_cat = default_cat.strip()
            cat_name, is_cat_valid = self.resolve_category(raw_cat, odoo_cats, default_fallback=default_cat)
            if not is_cat_valid:
                row_errors.append(f"التصنيف غير معرّف في أودو: '{raw_cat}'")

            # C. Tags (Brand tag is ALWAYS uppercase)
            tags_list = []
            if default_tag:
                tags_list.append(default_tag)
            if item['brand']:
                brand_tag = str(item['brand']).strip().upper()
                if brand_tag:
                    tags_list.append(brand_tag)
            if supplier_tag:
                tags_list.append(supplier_tag)

            clean_tags = [str(t).strip() for t in tags_list if str(t).strip()]
            tags_str = ",".join(clean_tags)

            # Check tags in Odoo
            if has_odoo:
                for t in tags_list:
                    t_clean = str(t).strip().upper()
                    if t_clean and t_clean.lower() not in odoo_tags:
                        row_warnings.append(f"الوسم غير مسجل في أودو: '{t_clean}'")
                        # For new items, record the missing tag for the 'التاجات الواجب تعريفها' sheet
                        if not is_bc_duplicate:
                            if t_clean not in missing_new_tags:
                                missing_new_tags[t_clean] = {
                                    'name': t_clean,
                                    'count': 0,
                                    'sample_barcodes': []
                                }
                            missing_new_tags[t_clean]['count'] += 1
                            if len(missing_new_tags[t_clean]['sample_barcodes']) < 3 and bc:
                                missing_new_tags[t_clean]['sample_barcodes'].append(bc)

            # D. Internal Reference
            # Rule: If barcode is duplicate in Odoo, DO NOT assign an Internal Reference
            # to preserve continuous sequence numbering for genuine new items.
            if is_bc_duplicate:
                internal_ref = ""
            else:
                prefix = self.find_internal_ref_prefix(cat_name or raw_cat, default_tag)
                internal_ref = get_next_seq_code(prefix)


            # E. Tax Validation
            tax_val = item['tax_raw']
            is_taxable, is_tax_valid = parse_tax_status(tax_val, self.config.get("tax_rules"))
            tax_err = not is_tax_valid

            if tax_err:
                row_errors.append(f"قيمة الضريبة غير معروفة: '{tax_val}'")

            # F. Sales Price & SO Price Calculation
            cost = item['ea_cost']
            sales_price = calculate_sales_price(cost, margin, is_taxable=is_taxable, tax_rate=0.15)
            so_price = sales_price

            # G. Names Formatting (trim and uppercase for EN)
            name_en = item['name_en'].strip().upper()
            name_ar = item['name_ar'].strip()

            # H. Customer & Vendor Taxes Database IDs
            if tax_err:
                row_cust_tax = ""
                row_vend_tax = ""
            elif is_taxable:
                row_cust_tax = customer_tax_id
                row_vend_tax = vendor_tax_id
            else:
                row_cust_tax = customer_tax_id_zero
                row_vend_tax = vendor_tax_id_zero

            processed_item = {
                'sequence': fixed_sequence,
                'internal_ref': internal_ref,
                'barcode': bc,
                'name': name_en,
                'name_ar': name_ar,
                'cost': cost,
                'sales_price': sales_price,
                'so_price': so_price,
                'category': cat_name or item['category_raw'],
                'tags': tags_str,
                'uom_po_id': uom_po_id,
                'uom_so_id': uom_so_id,
                'customer_tax_id': row_cust_tax,
                'vendor_tax_id': row_vend_tax,
                'type': prod_type,
                'invoice_policy': invoice_policy,
                'available_in_pos': available_in_pos,
                'errors': row_errors,
                'warnings': row_warnings,
                'is_bc_valid_chk': is_bc_valid_chk,
                'is_bc_duplicate': is_bc_duplicate,
                'is_cat_valid': is_cat_valid,
                'tax_err': tax_err,
                'raw_item': item
            }
            processed_rows.append(processed_item)

            if row_errors:
                total_errors += 1

            # Duplicate record for comparison sheet
            if is_bc_duplicate and bc in existing_products:
                ex = existing_products[bc]
                ex_cost = ex.get('cost', 0.0)
                ex_price = ex.get('price', 0.0)
                qty_on_hand = ex.get('qty_available', 0.0)
                cost_diff = cost - ex_cost
                price_diff = sales_price - ex_price

                # Determine difference notes
                diff_notes = []
                has_revaluation_risk = False

                if abs(cost_diff) > 0.01:
                    diff_notes.append(f"تغيير التكلفة: {cost_diff:+.2f} ر.س")
                    if qty_on_hand > 0:
                        has_revaluation_risk = True
                        diff_notes.append(f"⚠️ رصيد بالمخزون ({qty_on_hand:g}) - تغيير التكلفة يسبب Revaluation")

                if abs(price_diff) > 0.01:
                    diff_notes.append(f"تغيير سعر البيع: {price_diff:+.2f} ر.س")

                if ex.get('name_en', '').strip().upper() != name_en:
                    diff_notes.append("اختلاف الاسم الإنجليزي")

                if ex.get('name_ar', '').strip() != name_ar and name_ar:
                    diff_notes.append("اختلاف الاسم العربي")

                odoo_cat = (ex.get('category') or '').strip().lower()
                supp_cat = (cat_name or item['category_raw'] or '').strip().lower()
                if odoo_cat and supp_cat and odoo_cat != supp_cat:
                    diff_notes.append("اختلاف التصنيف")

                # ONLY include rows in comparison sheet if they have actual differences/notes!
                if diff_notes:
                    duplicate_rows.append({
                        'odoo_id': ex.get('odoo_id', ''),
                        'barcode': bc,
                        'odoo_ref': ex.get('default_code', ''),
                        'qty_on_hand': qty_on_hand,
                        'has_revaluation_risk': has_revaluation_risk,
                        'odoo_name_en': ex.get('name_en', ''),
                        'supp_name_en': name_en,
                        'odoo_name_ar': ex.get('name_ar', ''),
                        'supp_name_ar': name_ar,
                        'odoo_cost': ex_cost,
                        'supp_cost': cost,
                        'cost_diff': cost_diff,
                        'odoo_price': ex_price,
                        'supp_price': sales_price,
                        'price_diff': price_diff,
                        'odoo_category': ex.get('category', ''),
                        'supp_category': cat_name or item['category_raw'],
                        'notes': " | ".join(diff_notes),
                        'created_by': ex.get('created_by', ''),
                        'create_date': ex.get('create_date', '')
                    })

        notify(75, "جاري بناء ملف الإكسل Odooimport.xlsx وتنسيق الجداول...")

        # 4. Generate Odooimport.xlsx
        wb_out = openpyxl.Workbook()

        # Sheet 1: OdooImport
        ws_import = wb_out.active
        ws_import.title = "OdooImport"
        ws_import.views.sheetView[0].rightToLeft = False  # Odoo import template is LTR

        # Styles
        header_fill = PatternFill(start_color="1E6EBA", end_color="1E6EBA", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        error_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        error_font = Font(name="Calibri", size=11, color="9C0006")
        warn_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        warn_font = Font(name="Calibri", size=11, color="9C6500")

        thin_side = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        headers = [
            "Sequence",
            "Internal Reference",
            "Barcode",
            "Name",
            "Product Name Arabic",
            "Cost",
            "Sales Price",
            "SO Price",
            "Product Category",
            "Tags",
            "uom_po_id",
            "uom_so_id",
            "Customer Taxes / Database ID",
            "Vendor Taxes / Database ID",
            "type",
            "Invoicing Policy",
            "Available in POS",
            "seller_ids/partner_id/.id",
            "seller_ids/price",
            "seller_ids/min_qty",
            "seller_ids/delay",
            "ملاحظات التدقيق (Validation Notes)",
            "Description"
        ]

        ws_import.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws_import.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_idx, r in enumerate(processed_rows, start=2):
            all_notes = r['errors'] + r['warnings']
            notes_str = " | ".join(all_notes) if all_notes else ""

            vend_id_val = partner_id if partner_id else ""
            has_valid_supp = bool(vend_id_val and not r['is_bc_duplicate'] and r['internal_ref'])

            row_data = [
                r['sequence'],
                r['internal_ref'],
                r['barcode'],
                r['name'],
                r['name_ar'],
                r['cost'],
                r['sales_price'],
                r['so_price'],
                r['category'],
                r['tags'],
                r['uom_po_id'],
                r['uom_so_id'],
                r['customer_tax_id'],
                r['vendor_tax_id'],
                r['type'],
                r.get('invoice_policy', 'Delivered quantities'),
                "TRUE" if r['available_in_pos'] else "FALSE",
                vend_id_val if has_valid_supp else "",
                r['cost'] if has_valid_supp else "",
                1.0 if has_valid_supp else "",
                0 if has_valid_supp else "",
                notes_str,
                ""
            ]
            ws_import.append(row_data)

            # Apply Borders and Alignments
            for col_idx in range(1, len(row_data) + 1):
                c = ws_import.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                c.font = Font(name="Calibri", size=11)
                c.alignment = Alignment(vertical="center")

            # Format Numbers
            ws_import.cell(row=row_idx, column=6).number_format = '#,##0.00'
            ws_import.cell(row=row_idx, column=7).number_format = '#,##0.00'
            ws_import.cell(row=row_idx, column=8).number_format = '#,##0.00'
            ws_import.cell(row=row_idx, column=19).number_format = '#,##0.00'

            # Highlight specific error cells
            c_bc = ws_import.cell(row=row_idx, column=3)
            if not r['is_bc_valid_chk']:
                c_bc.fill = error_fill
                c_bc.font = error_font
            elif r['is_bc_duplicate']:
                c_bc.fill = warn_fill
                c_bc.font = warn_font

            c_cat = ws_import.cell(row=row_idx, column=9)
            if not r['is_cat_valid']:
                c_cat.fill = error_fill
                c_cat.font = error_font

            if r['tax_err']:
                ws_import.cell(row=row_idx, column=13).fill = error_fill
                ws_import.cell(row=row_idx, column=14).fill = error_fill

            # Validation note column styling (ملاحظات التدقيق)
            col_notes_idx = headers.index("ملاحظات التدقيق (Validation Notes)") + 1
            c_note = ws_import.cell(row=row_idx, column=col_notes_idx)
            if r['errors']:
                c_note.fill = error_fill
                c_note.font = error_font
            elif r['warnings']:
                c_note.fill = warn_fill
                c_note.font = warn_font

        # Auto-fit columns in Sheet 1
        for col in ws_import.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str))
            ws_import.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # 4.5 Add Excel Dropdown Data Validation for Column I (Product Category)
        category_options = self.get_odoo_category_names()
        if not category_options:
            category_options = list(DEFAULT_ODOO_CATEGORIES)

        ws_lookup = wb_out.create_sheet(title="CategoriesLookup")
        ws_lookup.views.sheetView[0].rightToLeft = False
        for c_idx, cat_text in enumerate(category_options, start=1):
            ws_lookup.cell(row=c_idx, column=1, value=cat_text)
        ws_lookup.sheet_state = "hidden"

        # Apply DataValidation to Column I (Product Category)
        last_row = max(len(processed_rows) + 50, 2000)
        dv_cat = DataValidation(
            type="list",
            formula1=f"=CategoriesLookup!$A$1:$A${len(category_options)}",
            allow_blank=True
        )
        dv_cat.error = "يرجى اختيار تصنيف صالح من قائمة تصنيفات أودو المنسدلة"
        dv_cat.errorTitle = "تصنيف غير معتمد"
        dv_cat.prompt = "اختر التصنيف المعتمد بأودو من القائمة المنسدلة"
        dv_cat.promptTitle = "تصنيف أودو"
        dv_cat.showInputMessage = False
        dv_cat.showErrorMessage = False

        ws_import.add_data_validation(dv_cat)
        dv_cat.add(f"I2:I{last_row}")

        # 5. Sheet 2: Duplicate Comparison Sheet (Only rows with differences/notes)
        if duplicate_rows:
            notify(88, f"تم رصد {len(duplicate_rows)} صنف مسجل مسبقاً به ملاحظات/فروقات. جاري بناء صفحة المقارنة...")
            ws_dup = wb_out.create_sheet(title="مقارنة الأصناف المسجلة")
            ws_dup.views.sheetView[0].rightToLeft = True  # Comparison sheet is RTL

            dup_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            dup_headers = [
                "id",
                "الباركود",
                "كود أودو الحالي",
                "الرصيد بالمخزون (Qty on Hand)",
                "الاسم إنجليزي (أودو)",
                "الاسم إنجليزي (المورد)",
                "الاسم عربي (أودو)",
                "الاسم عربي (المورد)",
                "التكلفة بأودو",
                "التكلفة الجديدة",
                "فرق التكلفة (+/-)",
                "سعر البيع بأودو",
                "السعر الجديد المقترح",
                "فرق سعر البيع (+/-)",
                "التصنيف بأودو",
                "التصنيف من المورد",
                "ملاحظات التغيير والفروقات",
                "من عرّف الصنف",
                "تاريخ التعريف"
            ]

            ws_dup.append(dup_headers)
            for col_idx in range(1, len(dup_headers) + 1):
                cell = ws_dup.cell(row=1, column=col_idx)
                cell.fill = dup_header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for d_idx, d in enumerate(duplicate_rows, start=2):
                d_row = [
                    d['odoo_id'],
                    d['barcode'],
                    d['odoo_ref'],
                    d['qty_on_hand'],
                    d['odoo_name_en'],
                    d['supp_name_en'],
                    d['odoo_name_ar'],
                    d['supp_name_ar'],
                    d['odoo_cost'],
                    d['supp_cost'],
                    d['cost_diff'],
                    d['odoo_price'],
                    d['supp_price'],
                    d['price_diff'],
                    d['odoo_category'],
                    d['supp_category'],
                    d['notes'],
                    d['created_by'],
                    d['create_date']
                ]
                ws_dup.append(d_row)

                for col_idx in range(1, len(d_row) + 1):
                    c = ws_dup.cell(row=d_idx, column=col_idx)
                    c.border = thin_border
                    c.font = Font(name="Calibri", size=11)
                    c.alignment = Alignment(vertical="center")

                # Column 1 (id): Center
                ws_dup.cell(row=d_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
                # Column 2 (barcode): Center
                ws_dup.cell(row=d_idx, column=2).alignment = Alignment(horizontal="center", vertical="center")
                # Column 3 (default_code): Center
                ws_dup.cell(row=d_idx, column=3).alignment = Alignment(horizontal="center", vertical="center")

                # Column 4 (Qty on Hand): Number format & Revaluation Warning
                c_qty = ws_dup.cell(row=d_idx, column=4)
                c_qty.number_format = '#,##0.00'
                c_qty.alignment = Alignment(horizontal="center", vertical="center")
                if d['has_revaluation_risk']:
                    c_qty.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
                    c_qty.font = Font(name="Calibri", size=11, bold=True, color="B25E00")
                elif d['qty_on_hand'] > 0:
                    c_qty.font = Font(name="Calibri", size=11, bold=True)

                # Format Numbers
                ws_dup.cell(row=d_idx, column=9).number_format = '#,##0.00'
                ws_dup.cell(row=d_idx, column=10).number_format = '#,##0.00'
                ws_dup.cell(row=d_idx, column=11).number_format = '+#,##0.00;-#,##0.00;0.00'
                ws_dup.cell(row=d_idx, column=12).number_format = '#,##0.00'
                ws_dup.cell(row=d_idx, column=13).number_format = '#,##0.00'
                ws_dup.cell(row=d_idx, column=14).number_format = '+#,##0.00;-#,##0.00;0.00'

                # Diff highlight
                c_cost_diff = ws_dup.cell(row=d_idx, column=11)
                if d['cost_diff'] > 0.01:
                    c_cost_diff.font = Font(name="Calibri", size=11, bold=True, color="9C0006")
                    c_cost_diff.fill = error_fill
                elif d['cost_diff'] < -0.01:
                    c_cost_diff.font = Font(name="Calibri", size=11, bold=True, color="006100")
                    c_cost_diff.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

                # Notes highlight if revaluation risk
                if d['has_revaluation_risk']:
                    c_notes = ws_dup.cell(row=d_idx, column=17)
                    c_notes.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
                    c_notes.font = Font(name="Calibri", size=11, bold=True, color="B25E00")

                # Column 18 (Created By): Center
                ws_dup.cell(row=d_idx, column=18).alignment = Alignment(horizontal="center", vertical="center")
                # Column 19 (Create Date): Center
                ws_dup.cell(row=d_idx, column=19).alignment = Alignment(horizontal="center", vertical="center")

            # Auto-fit columns in Sheet 2
            for col in ws_dup.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    max_len = max(max_len, len(val_str))
                ws_dup.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

        # 6. Sheet 3: Missing Tags Sheet (Only if there are undefined tags for new items)
        if missing_new_tags:
            notify(93, f"تم رصد {len(missing_new_tags)} وسم (Tag) غير معرّف لأصناف جديدة. جاري إنشاء ورقة التاجات...")
            ws_tags = wb_out.create_sheet(title="التاجات الواجب تعريفها")
            ws_tags.views.sheetView[0].rightToLeft = True

            tag_header_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")
            tag_headers = ["name", "عدد الأصناف المرتبطة", "أمثلة لباركودات الأصناف"]
            ws_tags.append(tag_headers)

            for col_idx in range(1, len(tag_headers) + 1):
                cell = ws_tags.cell(row=1, column=col_idx)
                cell.fill = tag_header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for t_idx, (t_name, t_info) in enumerate(missing_new_tags.items(), start=2):
                samples_str = ", ".join(t_info['sample_barcodes'])
                ws_tags.append([t_name, t_info['count'], samples_str])

                ws_tags.cell(row=t_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
                ws_tags.cell(row=t_idx, column=1).font = Font(name="Calibri", size=11, bold=True)
                ws_tags.cell(row=t_idx, column=1).border = thin_border

                ws_tags.cell(row=t_idx, column=2).alignment = Alignment(horizontal="center", vertical="center")
                ws_tags.cell(row=t_idx, column=2).font = Font(name="Calibri", size=11)
                ws_tags.cell(row=t_idx, column=2).border = thin_border

                ws_tags.cell(row=t_idx, column=3).alignment = Alignment(horizontal="center", vertical="center")
                ws_tags.cell(row=t_idx, column=3).font = Font(name="Calibri", size=11)
                ws_tags.cell(row=t_idx, column=3).border = thin_border

            for col in ws_tags.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    max_len = max(max_len, len(val_str))
                ws_tags.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 45)

        # 7. Sheet 4: Vendor Pricelist Sheet (Formatted for Odoo product.supplierinfo import)
        new_valid_items = [r for r in processed_rows if not r.get('is_bc_duplicate') and r.get('internal_ref')]
        if new_valid_items:
            notify(96, "جاري بناء ورقة قائمة أسعار المورد (Vendor Pricelist)...")
            ws_vendor = wb_out.create_sheet(title="قائمة أسعار المورد")
            ws_vendor.views.sheetView[0].rightToLeft = True

            v_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            v_headers = [
                "كود الصنف (product_tmpl_id/default_code)",
                "الباركود (Barcode)",
                "اسم الصنف (Name)",
                "معرف المورد بأودو (partner_id/.id)",
                "اسم المورد (partner_id)",
                "سعر شراء المورد (price)",
                "الكمية الدنيا (min_qty)",
                "مهلة التوريد بالأيام (delay)"
            ]
            ws_vendor.append(v_headers)
            for col_idx in range(1, len(v_headers) + 1):
                cell = ws_vendor.cell(row=1, column=col_idx)
                cell.fill = v_header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            vend_id_val = partner_id if partner_id else ""
            for v_idx, itm in enumerate(new_valid_items, start=2):
                v_row = [
                    itm['internal_ref'],
                    itm['barcode'],
                    itm['name'],
                    vend_id_val,
                    supplier_tag,
                    itm['cost'],
                    1.0,
                    0
                ]
                ws_vendor.append(v_row)
                for col_idx in range(1, len(v_row) + 1):
                    c = ws_vendor.cell(row=v_idx, column=col_idx)
                    c.border = thin_border
                    c.font = Font(name="Calibri", size=11)
                    c.alignment = Alignment(vertical="center")

                ws_vendor.cell(row=v_idx, column=1).alignment = Alignment(horizontal="center", vertical="center")
                ws_vendor.cell(row=v_idx, column=2).alignment = Alignment(horizontal="center", vertical="center")
                ws_vendor.cell(row=v_idx, column=4).alignment = Alignment(horizontal="center", vertical="center")
                ws_vendor.cell(row=v_idx, column=6).number_format = '#,##0.00'
                ws_vendor.cell(row=v_idx, column=7).alignment = Alignment(horizontal="center", vertical="center")
                ws_vendor.cell(row=v_idx, column=8).alignment = Alignment(horizontal="center", vertical="center")

            for col in ws_vendor.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    max_len = max(max_len, len(val_str))
                ws_vendor.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 45)

        # Ensure OdooImport is the active visible tab
        wb_out.active = ws_import

        # Save Odooimport.xlsx
        saved_file = output_import_path
        save_warning = ""
        try:
            wb_out.save(output_import_path)
        except PermissionError:
            # Fallback if file is locked or open in Excel
            base, ext = os.path.splitext(output_import_path)
            saved_file = f"{base}_جديد{ext}"
            try:
                wb_out.save(saved_file)
                save_warning = f"تم الحفظ باسم '{os.path.basename(saved_file)}' لأن الملف الأصلي مفتوح في برنامج إكسل."
            except Exception as e2:
                raise PermissionError(f"تعذر حفظ الملف لأن الملف مفتوح حالياً في برنامج إكسل (Excel): {e2}")
        except Exception as e:
            raise RuntimeError(f"خطأ أثناء حفظ ملف الإكسل: {e}")

        if save_warning:
            notify(100, f"✅ اكتمل تجهيز الملف: {os.path.basename(saved_file)} ({save_warning})")
        else:
            notify(100, f"✅ اكتمل تجهيز الملف: {os.path.basename(saved_file)}")

        total_duplicates = sum(1 for r in processed_rows if r.get('is_bc_duplicate'))
        return {
            'total_items': len(data_rows),
            'total_errors': total_errors,
            'duplicate_count': len(duplicate_rows),
            'total_duplicates': total_duplicates,
            'missing_tags_count': len(missing_new_tags),
            'missing_tags': list(missing_new_tags.keys()),
            'supplier_no': supplier_no,
            'supplier_name': supplier_name,
            'partner_id': partner_id,
            'vendor_info': vendor_info,
            'output_file': saved_file,
            'save_warning': save_warning,
            'can_archive': total_errors == 0,
            'processed_rows': processed_rows,
            'duplicate_rows': duplicate_rows
        }

    def append_to_all_new_items(
        self,
        supp_items_path: str,
        all_new_path: str,
        vendor_override: Optional[Dict[str, Any]] = None,
        category_override: Optional[str] = None
    ) -> int:
        """
        Appends valid items from supplier file into AllNewItems.xlsx,
        mapping columns dynamically and adding Supplier No and Supplier Name.
        """
        wb_supp = openpyxl.load_workbook(supp_items_path, data_only=True)
        ws_supp = wb_supp.active

        # Extract Supplier No and Name
        if vendor_override and vendor_override.get('supplier_no'):
            supplier_no = vendor_override.get('supplier_no')
            supplier_name = vendor_override.get('supplier_name', '')
        else:
            hint = extract_vendor_hint(supp_items_path, ws_supp)
            query = hint.get('code') or hint.get('name') or ""
            if query:
                v_info = self.verifier.resolve_vendor(query)
            else:
                v_info = {'success': False}
            supplier_no = v_info.get('supplier_no') or hint.get('code') or "SUPP"
            supplier_name = v_info.get('supplier_name') or hint.get('name') or "Supplier"

        header_row_idx, col_map, _ = detect_header_row_and_mapping(ws_supp)

        wb_all = openpyxl.load_workbook(all_new_path)
        ws_all = wb_all.active

        # Check existing barcodes in AllNewItems to prevent duplicate appending
        existing_all_barcodes = set()
        for r in range(3, ws_all.max_row + 1):
            bc = ws_all.cell(row=r, column=4).value
            if bc:
                existing_all_barcodes.add(str(bc).strip())

        STANDARD_FIELDS = [
            'brand', 'name_en', 'name_ar', 'ea_barcode', 'cs_barcode', 'cs_size',
            'cs_cost', 'ea_cost', 'origin', 'picture', 'shelf_life', 'tax',
            'category', 'sub_category'
        ]

        appended_count = 0
        for r in range(header_row_idx + 1, ws_supp.max_row + 1):
            row_cells = ws_supp[r]
            row_vals = [cell.value for cell in row_cells]
            if not any(v is not None for v in row_vals):
                continue

            def get_val(f_name: str, default: Any = None):
                idx = col_map.get(f_name)
                if idx is not None and idx < len(row_vals):
                    v = row_vals[idx]
                    return v if v is not None else default
                return default

            ea_barcode = str(get_val('ea_barcode', '') or '').strip()
            if ea_barcode and ea_barcode in existing_all_barcodes:
                continue

            standard_row = [get_val(f) for f in STANDARD_FIELDS]
            # Brand uppercase
            if standard_row[0]:
                standard_row[0] = str(standard_row[0]).strip().upper()
            if standard_row[6] is not None:
                standard_row[6] = safe_float(standard_row[6])
            if standard_row[7] is not None:
                standard_row[7] = safe_float(standard_row[7])
            # Fallback category if blank
            if (not standard_row[12] or not str(standard_row[12]).strip()) and category_override:
                standard_row[12] = category_override.strip()


            full_row = standard_row + [supplier_no, supplier_name]
            ws_all.append(full_row)
            appended_count += 1
            if ea_barcode:
                existing_all_barcodes.add(ea_barcode)

        wb_all.save(all_new_path)
        return appended_count

